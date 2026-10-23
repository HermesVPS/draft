"""Contact-list file parsing (.xlsx and .csv), ported from the MVP.

The MVP used the JS ``xlsx`` library for spreadsheets; we use openpyxl. The
CSV path reuses ``validation.parse_csv``.
"""

import io

from openpyxl import load_workbook

from .validation import parse_csv

# Temporary Excel lock files (e.g. "~$contacts.xlsx") must be ignored.
LOCK_FILE_PREFIX = "~$"


def is_lock_file(filename):
    return str(filename or "").startswith(LOCK_FILE_PREFIX)


def parse_xlsx_bytes(data, sheet_name=None):
    """Parse an .xlsx byte stream into ``{headers, rows, sheet_names}``.

    Reads the first worksheet by default (or ``sheet_name`` when given).
    Preserves cell values as text and drops fully-empty rows.
    """
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        if not sheet_names:
            return {"headers": [], "rows": [], "sheet_names": []}
        ws = workbook[sheet_name] if sheet_name and sheet_name in sheet_names else workbook[sheet_names[0]]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return {"headers": [], "rows": [], "sheet_names": sheet_names}
        headers = [("" if c is None else str(c).strip()) for c in header_row]
        rows = []
        for raw in rows_iter:
            values = [("" if c is None else str(c).strip()) for c in raw]
            if any(values):
                rows.append(dict(zip(headers, values)))
        return {"headers": headers, "rows": rows, "sheet_names": sheet_names}
    finally:
        workbook.close()


def parse_contacts_file(filename, data, sheet_name=None):
    """Dispatch to the right parser based on the file extension.

    Always returns ``{headers, rows, sheet_names}`` (``sheet_names`` is empty
    for CSV).
    """
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        result = parse_xlsx_bytes(data, sheet_name=sheet_name)
    elif name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace") if isinstance(data, (bytes, bytearray)) else str(data)
        result = parse_csv(text)
    else:
        raise ValueError("Unsupported contact list type. Upload a .xlsx or .csv file.")
    result.setdefault("sheet_names", [])
    return result
